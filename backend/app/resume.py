"""Resume upload: extract text from PDF / DOCX / TXT and summarise it into a candidate profile."""

import io
import json
import logging

from pydantic import BaseModel, Field, ValidationError

from app.groq_client import AIError, chat

logger = logging.getLogger("interview_coach")

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_RESUME_CHARS = 8000  # keeps the prompt small; a typical 2-page resume is ~4-5k chars

ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "txt",
}
ALLOWED_EXTENSIONS = {"pdf": "pdf", "docx": "docx", "txt": "txt", "md": "txt"}


class ResumeProfile(BaseModel):
    summary: str = Field(min_length=1)
    experience_level: str = Field(min_length=1)
    key_skills: list[str] = Field(max_length=12)
    notable_projects: list[str] = Field(max_length=6)


def detect_kind(filename: str, content_type: str | None) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ALLOWED_EXTENSIONS:
        return ALLOWED_EXTENSIONS[ext]
    if content_type in ALLOWED_TYPES:
        return ALLOWED_TYPES[content_type]
    raise AIError("Unsupported file type. Please upload a PDF, Word (.docx), or plain-text resume.", status=422)


def extract_text(data: bytes, kind: str) -> str:
    try:
        if kind == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif kind == "docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="replace")
    except Exception as e:  # corrupt / encrypted / not really that format
        logger.warning("Resume extraction failed: %s", e)
        raise AIError("We couldn't read that file. Make sure it isn't password-protected or corrupted.", status=422)

    # Collapse whitespace but keep line structure.
    lines = [" ".join(line.split()) for line in text.splitlines()]
    text = "\n".join(line for line in lines if line)
    if len(text) < 80:
        raise AIError(
            "We couldn't find enough text in that file. If it's a scanned PDF, try exporting a text-based version.",
            status=422,
        )
    return text[:MAX_RESUME_CHARS]


PROFILE_PROMPT = """You summarise a candidate's resume for a technical interviewer.
Return ONLY a JSON object:
{
  "summary": "<2 sentences: who this candidate is and what they have done, in plain language>",
  "experience_level": "<one of: Student/Intern, Junior, Mid-level, Senior, Staff/Lead, Manager>",
  "key_skills": [<up to 10 concrete technologies or skills that appear in the resume>],
  "notable_projects": [<up to 4 short phrases naming specific projects, products or achievements from the resume>]
}
Only use information present in the resume. Do not invent anything."""


def build_profile(resume_text: str) -> ResumeProfile | None:
    """Ask the model for a short profile. Returns None if the AI is unavailable — the upload still succeeds."""
    try:
        raw = chat(
            [
                {"role": "system", "content": PROFILE_PROMPT},
                {"role": "user", "content": f"RESUME:\n\n{resume_text}"},
            ],
            temperature=0.2,
            json_mode=True,
            max_tokens=1500,
        )
        return ResumeProfile.model_validate(json.loads(raw))
    except (AIError, json.JSONDecodeError, ValidationError) as e:
        logger.warning("Resume profile generation failed: %s", e)
        return None
